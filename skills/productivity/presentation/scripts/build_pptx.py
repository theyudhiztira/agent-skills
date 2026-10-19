#!/usr/bin/env python3
"""Build a presentation-ready .pptx from a deck spec (JSON) and a token preset.

Usage
  python3 build_pptx.py deck.json -o out.pptx            # tokens from deck.meta.preset
  python3 build_pptx.py deck.json -t ../tokens/neutral-light.json -o out.pptx
  python3 build_pptx.py deck.json --check                 # lint only, no file

The deck spec format is documented in references/deck-spec.md.
Every layout name maps to one function below (LAYOUTS at the bottom).
Coordinates are written on a 1920 x 1080 px canvas and converted to EMU.

Requires: python-pptx (pip install python-pptx), Pillow for image cropping.
"""
import argparse
import json
import math
import os
import re
import sys

try:
    from pptx import Presentation
    from pptx.util import Emu, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
    from pptx.enum.dml import MSO_LINE
    from pptx.oxml.ns import qn
    from lxml import etree
except ImportError:  # pragma: no cover
    sys.exit("python-pptx is missing. Install it with: pip install python-pptx")

HERE = os.path.dirname(os.path.abspath(__file__))
TOKENS_DIR = os.path.normpath(os.path.join(HERE, "..", "tokens"))

EMU_PER_PX = 6350          # 13.333 in / 1920 px
CANVAS_W, CANVAS_H = 1920, 1080

WARNINGS = []


def warn(slide_no, msg):
    WARNINGS.append(f"slide {slide_no:02d}: {msg}")


# --------------------------------------------------------------------------- units

def px(v):
    return Emu(int(round(v * EMU_PER_PX)))


def fpt(size_px):
    """Font size: 1 canvas px = 0.5 pt on a 13.333 in wide slide."""
    return Pt(size_px * 0.5)


def rgb(hexstr):
    h = hexstr.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


# --------------------------------------------------------------------------- tokens

class Theme:
    def __init__(self, data):
        self.d = data
        c = data["color"]
        self.c = c
        self.t = data["type"]
        self.s = self.t["scale"]
        self.l = data["layout"]
        self.decor = data.get("decor", {})

    def col(self, key):
        """Resolve a colour token name (dotted for nested) or a raw hex."""
        if isinstance(key, str) and key.startswith("#"):
            return key
        node = self.c
        for part in key.split("."):
            node = node[part]
        return node

    @property
    def sans(self):
        return self.t["fontSans"]

    @property
    def mono(self):
        return self.t.get("fontMono", self.t["fontSans"])


def load_tokens(path_or_name):
    path = path_or_name
    if not os.path.isfile(path):
        path = os.path.join(TOKENS_DIR, f"{path_or_name}.json")
    with open(path, encoding="utf-8") as fh:
        return Theme(json.load(fh))


# --------------------------------------------------------------------------- text measuring

def est_lines(text, size, width, upper=False, mono=False, tracking=0):
    """Rough wrap estimate. Good enough to warn about overflow, not to typeset."""
    factor = 0.6 if mono else (0.66 if upper else 0.53)
    if size >= 96 and not mono:
        factor -= 0.05  # display sizes are set light and tight
    char_w = size * factor + tracking
    per_line = max(1, int(width / char_w))
    lines = 0
    for para in str(text).split("\n"):
        words = para.split(" ")
        cur = 0
        n = 1
        for w in words:
            wl = len(w)
            if cur == 0:
                cur = wl
            elif cur + 1 + wl <= per_line:
                cur += 1 + wl
            else:
                n += 1
                cur = wl
        lines += n
    return lines


def est_height(text, size, width, lh=1.3, **kw):
    return est_lines(text, size, width, **kw) * size * lh


# --------------------------------------------------------------------------- drawing primitives

INLINE = re.compile(r"(\*\*.+?\*\*|==.+?==)")


def _style_run(run, theme, size, color, bold=False, italic=False, font=None, tracking=0):
    f = run.font
    f.name = font or theme.sans
    f.size = fpt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = rgb(theme.col(color))
    if tracking:
        run._r.get_or_add_rPr().set("spc", str(int(tracking * 0.5 * 100)))


def text(slide, theme, x, y, w, h, content, size, color="ink", bold=False, italic=False,
         align="left", lh=1.25, font=None, upper=False, tracking=0, anchor="top",
         para_space=0, slide_no=0, check=True):
    """Add a text box. `content` is a string (\\n = new paragraph) or a list of
    paragraphs. Inline markup: **bold**, ==accent colour==."""
    tb = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE,
                          "bottom": MSO_ANCHOR.BOTTOM}[anchor]
    paras = content if isinstance(content, list) else str(content).split("\n")
    for i, ptxt in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
                       "right": PP_ALIGN.RIGHT}[align]
        p.line_spacing = lh
        if para_space and i > 0:
            p.space_before = fpt(para_space)
        ptxt = ptxt.upper() if upper else ptxt
        for seg in INLINE.split(ptxt):
            if not seg:
                continue
            r = p.add_run()
            if seg.startswith("**") and seg.endswith("**"):
                r.text = seg[2:-2]
                _style_run(r, theme, size, color, True, italic, font, tracking)
            elif seg.startswith("==") and seg.endswith("=="):
                r.text = seg[2:-2]
                _style_run(r, theme, size, "accentText", bold, italic, font, tracking)
            else:
                r.text = seg
                _style_run(r, theme, size, color, bold, italic, font, tracking)
    if check:
        plain = INLINE.sub(lambda m: m.group(0)[2:-2], "\n".join(paras))
        need = est_height(plain, size, w, lh, upper=upper, mono=(font == theme.mono),
                          tracking=tracking)
        if need > h * 1.08:
            warn(slide_no, f"text may overflow its box ({int(need)}px needed, {int(h)}px given): "
                           f"\"{plain[:48]}...\"")
    return tb


def rect(slide, x, y, w, h, fill=None, line=None, line_w=1, dash=False, alpha=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, px(x), px(y), px(w), px(h))
    shp.shadow.inherit = False
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
        if alpha is not None:
            clr = shp.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
            etree.SubElement(clr, qn("a:alpha")).set("val", str(int(alpha * 100000)))
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = px(line_w)
        if dash:
            shp.line.dash_style = MSO_LINE.DASH
    else:
        shp.line.fill.background()
    if shp.has_text_frame:
        shp.text_frame.text = ""
    return shp


def hline(slide, x1, x2, y, color, w=1, dash=False):
    return seg(slide, x1, y, x2, y, color, w, dash)


def vline(slide, x, y1, y2, color, w=1, dash=False):
    return seg(slide, x, y1, x, y2, color, w, dash)


def seg(slide, x1, y1, x2, y2, color, w=1, dash=False, arrow=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, px(x1), px(y1), px(x2), px(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = px(w)
    if dash:
        c.line.dash_style = MSO_LINE.DASH
    if arrow:
        ln = c.line._get_or_add_ln()
        etree.SubElement(ln, qn("a:tailEnd")).set("type", "triangle")
    return c


def curve(slide, p0, p1, p2, p3, color, w=2, dash=True, steps=24):
    """Cubic bezier drawn as a freeform polyline."""
    pts = []
    for i in range(steps + 1):
        t = i / steps
        x = (1 - t) ** 3 * p0[0] + 3 * (1 - t) ** 2 * t * p1[0] + 3 * (1 - t) * t ** 2 * p2[0] + t ** 3 * p3[0]
        y = (1 - t) ** 3 * p0[1] + 3 * (1 - t) ** 2 * t * p1[1] + 3 * (1 - t) * t ** 2 * p2[1] + t ** 3 * p3[1]
        pts.append((int(x * EMU_PER_PX), int(y * EMU_PER_PX)))
    b = slide.shapes.build_freeform(pts[0][0], pts[0][1], scale=1.0)
    b.add_line_segments(pts[1:], close=False)
    shp = b.convert_to_shape()
    shp.fill.background()
    shp.shadow.inherit = False
    shp.line.color.rgb = rgb(color)
    shp.line.width = px(w)
    if dash:
        shp.line.dash_style = MSO_LINE.DASH
    return shp


def image(slide, path, x, y, w, h, fit="cover"):
    if not path or not os.path.isfile(path):
        return None
    if fit == "contain":
        try:
            from PIL import Image
            iw, ih = Image.open(path).size
            r = min(w / iw, h / ih)
            nw, nh = iw * r, ih * r
            return slide.shapes.add_picture(path, px(x), px(y + (h - nh) / 2), px(nw), px(nh))
        except Exception:
            pass
    pic = slide.shapes.add_picture(path, px(x), px(y), px(w), px(h))
    if fit == "cover":
        try:
            from PIL import Image
            iw, ih = Image.open(path).size
            box_r, img_r = w / h, iw / ih
            if img_r > box_r:
                cut = (1 - box_r / img_r) / 2
                pic.crop_left = pic.crop_right = cut
            else:
                cut = (1 - img_r / box_r) / 2
                pic.crop_top = pic.crop_bottom = cut
        except Exception:
            pass
    return pic


def square(slide, x, y, size, theme, style):
    """Status marker. style: filled | outline | muted."""
    if style == "filled":
        return rect(slide, x, y, size, size, fill=theme.col("accent"))
    if style == "outline":
        return rect(slide, x, y, size, size, line=theme.col("accentText"), line_w=2)
    return rect(slide, x, y, size, size, fill=theme.col("line"), line=theme.col("lineStrong"), line_w=2)


# --------------------------------------------------------------------------- slide chrome

def new_slide(prs, theme, ground="ground"):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(theme.col(ground))
    return s


def decor_grid(slide, theme, seed=0):
    d = theme.decor
    if not d.get("grid"):
        return
    pitch = d.get("gridPitch", 132)
    x0 = theme.l["marginX"]
    color = theme.col("line")
    x = x0
    while x < CANVAS_W:
        vline(slide, x, 0, CANVAS_H, color, 1)
        x += pitch
    y = 144
    while y < CANVAS_H:
        hline(slide, 0, CANVAS_W, y, color, 1)
        y += pitch
    # three nodes on intersections, placed from a seed so slides differ
    cols = int((CANVAS_W - x0) / pitch)
    rows = int((CANVAS_H - 144) / pitch)
    n = d.get("nodeSize", 14)
    for k in range(3):
        cx = x0 + ((seed * 5 + k * 7) % cols) * pitch
        cy = 144 + ((seed * 3 + k * 4) % rows) * pitch
        rect(slide, cx - n / 2, cy - n / 2, n, n, fill=theme.col("lineStrong"))


def chrome(slide, theme, meta, page=None, logos=True, dark_photo=False):
    m = theme.l["marginX"]
    if logos:
        if meta.get("logoLeft"):
            image(slide, meta["logoLeft"], m, 40, 200, 64, fit="contain")
        if meta.get("logoRight"):
            image(slide, meta["logoRight"], CANVAS_W - m - 220, 44, 220, 56, fit="contain")
    small = theme.s["small"]
    if meta.get("footer"):
        text(slide, theme, m, CANVAS_H - 64 - small * 1.3, 1300, small * 1.4, meta["footer"], small,
             color="inkMuted" if dark_photo else "inkSubtle", **lab(theme), check=False)
    if page is not None and meta.get("pageNumbers", True):
        text(slide, theme, CANVAS_W - m - 200, CANVAS_H - 64 - small * 1.3, 200, small * 1.4,
             f"{page:02d}", small, color="inkSubtle", align="right", font=theme.mono, check=False)



def lab(theme):
    """Small-label style. Uppercase + tracking only when the preset asks for it."""
    up = theme.decor.get("labelUppercase", False)
    return {"upper": up, "tracking": 3 if up else 0}

def eyebrow_style(theme):
    up = theme.decor.get("eyebrowUppercase", False)
    return dict(upper=up, tracking=4 if up else 0)


def header(slide, theme, sd, n):
    """Eyebrow + headline at the top margin. Returns the y where content starts."""
    m, top = theme.l["marginX"], theme.l["marginTop"]
    small, h1 = theme.s["small"], theme.s["h1"]
    y = top
    if sd.get("eyebrow"):
        text(slide, theme, m, y, 1728, small * 1.4, sd["eyebrow"], small, color="accentText",
             bold=True, slide_no=n, **eyebrow_style(theme))
        y += small * 1.4 + 12
    title = sd.get("title", "")
    th = est_height(title, h1, 1728, 1.15)
    text(slide, theme, m, y, 1728, th + 4, title, h1, color="ink", lh=1.15, slide_no=n, check=False)
    if len(title) > 70:
        warn(n, f"headline is long ({len(title)} chars); aim for one line, two at most")
    return y + th + 36


def label(slide, theme, x, y, w, s, color="inkMuted", n=0):
    small = theme.s["small"]
    text(slide, theme, x, y, w, small * 1.4, s, small, color=color, bold=True, **lab(theme), slide_no=n)
    return small * 1.4


def bottom_limit(theme):
    return CANVAS_H - theme.l["marginBottom"]


def check_bottom(theme, y, n, what="content"):
    if y > bottom_limit(theme) + 8:
        warn(n, f"{what} runs past the safe area by {int(y - bottom_limit(theme))}px; "
                f"split the slide or cut items")


# --------------------------------------------------------------------------- layouts

def L_cover(prs, theme, meta, sd, n, closing=False):
    s = new_slide(prs, theme, "night" if sd.get("image") else "ground")
    m = theme.l["marginX"]
    if sd.get("image"):
        image(s, sd["image"], 0, 0, CANVAS_W, CANVAS_H)
        rect(s, 0, 0, CANVAS_W, CANVAS_H, fill=theme.col("night"), alpha=0.72)
    decor_grid(s, theme, n)
    if meta.get("logoRight"):
        image(s, meta["logoRight"], CANVAS_W - m - 280, 40, 280, 72, fit="contain")
    on_photo = bool(sd.get("image"))
    ink = "inkInverse" if on_photo else "ink"
    disp = theme.s["display"]
    title = sd.get("title", "")
    upper = theme.decor.get("displayUppercase", False)
    th = est_height(title, disp, 1728, 1.0, upper=upper)
    lead = sd.get("subtitle", "")
    meta_line = sd.get("meta", "")
    lh_lead = theme.s["lead"] * 1.5
    block = th + (24 + est_height(lead, theme.s["lead"], 1500, 1.5) if lead else 0) + (16 + lh_lead if meta_line else 0)
    y = CANVAS_H - theme.l["marginBottom"] - block
    logo = meta.get("logoCover") or meta.get("logoLeft")
    if logo:
        image(s, logo, m, y - 141 - 40, 440, 141, fit="contain")
    text(s, theme, m, y, 1728, th + 8, title, disp, color=ink, lh=1.0, upper=upper, slide_no=n, check=False)
    y += th + 24
    if lead:
        h = est_height(lead, theme.s["lead"], 1500, 1.5)
        text(s, theme, m, y, 1500, h, lead, theme.s["lead"], color=ink, slide_no=n)
        y += h + 16
    if meta_line:
        text(s, theme, m, y, 1500, lh_lead, meta_line, theme.s["small"] + 2,
             color="inkMuted" if not on_photo else "inkInverse", slide_no=n)
    chrome(s, theme, meta, page=None, logos=False, dark_photo=on_photo)
    return s


def L_closing(prs, theme, meta, sd, n):
    return L_cover(prs, theme, meta, sd, n, closing=True)


def L_divider(prs, theme, meta, sd, n):
    s = new_slide(prs, theme)
    decor_grid(s, theme, n)
    m = theme.l["marginX"]
    idx = sd.get("index", "")
    if theme.decor.get("ghostNumeral", False) and idx:
        digit = str(int(re.sub(r"\D", "", idx) or 0))
        text(s, theme, 1240, 300, 640, 700, digit, 720, color="lineStrong", bold=True, align="right",
             lh=0.8, check=False)
    chrome(s, theme, meta, page=None)
    sec = theme.s["section"]
    upper = theme.decor.get("displayUppercase", False)
    title = sd.get("title", "")
    th = est_height(title, sec, 1500, 1.0, upper=upper)
    lead = sd.get("lead", "")
    lh = est_height(lead, theme.s["lead"], 1100, 1.5) if lead else 0
    total = theme.s["index"] + 24 + th + (24 + lh if lead else 0)
    y = (CANVAS_H - total) / 2
    if idx:
        text(s, theme, m, y, 600, theme.s["index"] * 1.05, f"{idx}_", theme.s["index"], color="accent",
             lh=1.0, check=False)
        y += theme.s["index"] + 24
    text(s, theme, m, y, 1500, th + 8, title, sec, color="ink", lh=1.0, upper=upper, slide_no=n)
    y += th + 24
    if lead:
        text(s, theme, m, y, 1100, lh, lead, theme.s["lead"], color="inkMuted", lh=1.5, slide_no=n)
        y += lh + 24
    tags = sd.get("tags", [])
    if tags:
        x = m
        for tg in tags:
            w = len(tg["label"]) * theme.s["small"] * 0.72 + 40
            color = theme.col(tg.get("color", "accentText"))
            rect(s, x, y, w, theme.s["small"] * 1.9, line=color, line_w=1)
            text(s, theme, x, y, w, theme.s["small"] * 1.9, tg["label"], theme.s["small"], color=color,
                 bold=True, align="center", anchor="middle", **lab(theme), check=False)
            x += w + 16
    return s


def L_summary(prs, theme, meta, sd, n):
    """Headline + lead + a ruled strip of 3-5 numbers."""
    s = new_slide(prs, theme)
    decor_grid(s, theme, n)
    chrome(s, theme, meta, page=n)
    m = theme.l["marginX"]
    y = header(s, theme, sd, n)
    if sd.get("lead"):
        h = est_height(sd["lead"], theme.s["lead"], 1400, 1.5)
        text(s, theme, m, y, 1400, h, sd["lead"], theme.s["lead"], color="inkMuted", lh=1.5, slide_no=n)
        y += h
    metrics = sd.get("metrics", [])
    if not 2 <= len(metrics) <= 5:
        warn(n, "summary works best with 3 to 5 numbers")
    mh = theme.s["metric"]
    strip_h = 32 + mh + 16 + theme.s["small"] * 1.4 + 10 + theme.s["small"] * 1.4 * 2 + 32
    top = bottom_limit(theme) - strip_h
    if top < y + 24:
        warn(n, "lead and numbers collide; shorten the lead")
    hline(s, m, m + 1728, top, theme.col("lineStrong"), 1)
    hline(s, m, m + 1728, top + strip_h, theme.col("line"), 1)
    cw = 1728 / max(1, len(metrics))
    for i, mt in enumerate(metrics):
        x = m + i * cw
        pad = 0 if i == 0 else 32
        if i:
            vline(s, x, top, top + strip_h, theme.col("line"))
        yy = top + 32
        text(s, theme, x + pad, yy, cw - pad - 24, mh * 1.05, str(mt["value"]), mh,
             color="accent" if i == 0 else "ink", font=theme.mono, lh=1.0, slide_no=n)
        yy += mh + 16
        label(s, theme, x + pad, yy, cw - pad - 24, mt.get("label", ""), color="ink", n=n)
        yy += theme.s["small"] * 1.4 + 10
        if mt.get("caption"):
            text(s, theme, x + pad, yy, cw - pad - 24, theme.s["small"] * 1.4 * 2, mt["caption"],
                 theme.s["small"], color="inkMuted", lh=1.4, slide_no=n)
    return s


def L_statement(prs, theme, meta, sd, n):
    """One big sentence. Use for the single most important message."""
    s = new_slide(prs, theme)
    decor_grid(s, theme, n)
    chrome(s, theme, meta, page=n)
    m = theme.l["marginX"]
    size = theme.s["statement"]
    st = sd.get("statement", "")
    h = est_height(st, size, 1500, 1.15)
    sub = sd.get("support", "")
    sh = est_height(sub, theme.s["lead"], 1300, 1.5) if sub else 0
    y = (CANVAS_H - (h + 40 + sh)) / 2
    if sd.get("eyebrow"):
        text(s, theme, m, y - 60, 1500, 40, sd["eyebrow"], theme.s["small"], color="accentText",
             bold=True, **eyebrow_style(theme), slide_no=n)
    text(s, theme, m, y, 1500, h, st, size, color="ink", lh=1.15, slide_no=n)
    if sub:
        text(s, theme, m, y + h + 40, 1300, sh, sub, theme.s["lead"], color="inkMuted", lh=1.5, slide_no=n)
    return s


def L_convergence(prs, theme, meta, sd, n):
    """Many separate inputs -> one engine -> one combined output."""
    s = new_slide(prs, theme)
    decor_grid(s, theme, n)
    chrome(s, theme, meta, page=n)
    m = theme.l["marginX"]
    y0 = header(s, theme, sd, n)
    inputs = sd.get("inputs", [])
    if not 2 <= len(inputs) <= 6:
        warn(n, "convergence needs 2 to 6 inputs")
    body = theme.s["body"]
    top = y0 + 8
    label(s, theme, m, top, 440, sd.get("inputsLabel", "Inputs"), color="inkSubtle", n=n)
    row_top = top + 48
    avail = bottom_limit(theme) - 110 - row_top
    step = min(110, avail / max(1, len(inputs)))
    centers = []
    for i, it in enumerate(inputs):
        cy = row_top + i * step + 31
        centers.append(cy)
        text(s, theme, m, cy - 31, 440, body * 1.25, it["label"], body, color="ink", slide_no=n)
        if it.get("sub"):
            text(s, theme, m, cy - 31 + body * 1.25, 440, body * 1.25, it["sub"], body,
                 color="inkMuted", slide_no=n)
        rect(s, 548, cy - 6, 12, 12, fill=theme.col("inkSubtle"))
    mid = (centers[0] + centers[-1]) / 2 if centers else row_top + 150
    ex, ew, eh = 900, 200, 140
    for cy in centers:
        curve(s, (560, cy), (760, cy), (740, mid), (ex, mid), theme.col("inkSubtle"), 2, dash=True)
    eng = sd.get("engine", {})
    rect(s, ex, mid - eh / 2, ew, eh, fill=theme.col("ground"), line=theme.col("accent"), line_w=2)
    text(s, theme, ex, mid - eh / 2, ew, eh, eng.get("label", "ENGINE").split("\n"), theme.s["small"],
         color="accentText", bold=True, align="center", anchor="middle", **lab(theme),
         slide_no=n)
    if eng.get("verbs"):
        vb = eng["verbs"]
        text(s, theme, ex - 50, mid + eh / 2 + 20, ew + 100, len(vb) * theme.s["small"] * 1.45,
             vb, theme.s["small"], color="inkMuted", align="center", lh=1.4, slide_no=n)
    seg(s, ex + ew, mid, 1250, mid, theme.col("accent"), 4, arrow=True)
    out = sd.get("output", {})
    ox, ow = 1296, 528
    oy = top
    label(s, theme, ox, oy, ow, out.get("label", "Output"), color="accentText", n=n)
    oy += 48
    if out.get("headline"):
        hh = est_height(out["headline"], theme.s["h2"], ow, 1.15)
        text(s, theme, ox, oy, ow, hh, out["headline"], theme.s["h2"], color="ink", lh=1.15, slide_no=n)
        oy += hh + 24
    for i, row in enumerate(out.get("rows", [])):
        hline(s, ox, ox + ow, oy, theme.col("lineStrong" if i == 0 else "line"))
        rh = est_height(row, theme.s["small"], ow, 1.3) + 28
        text(s, theme, ox, oy + 14, ow, rh - 20, row, theme.s["small"], color="ink", lh=1.3, slide_no=n)
        oy += rh
    if out.get("rows"):
        hline(s, ox, ox + ow, oy, theme.col("line"))
    if sd.get("takeaway"):
        ty = bottom_limit(theme) - 60
        text(s, theme, m, ty, 1728, 60, sd["takeaway"], theme.s["lead"], color="ink", lh=1.4, slide_no=n)
    check_bottom(theme, oy, n, "output column")
    return s


def L_matrix(prs, theme, meta, sd, n):
    """Columns = categories (2-4), rows = aspects (2-4)."""
    s = new_slide(prs, theme)
    decor_grid(s, theme, n)
    chrome(s, theme, meta, page=n)
    m = theme.l["marginX"]
    y = header(s, theme, sd, n) + 12
    cols = sd.get("columns", [])
    rows = sd.get("rows", [])
    lw = 240
    cw = (1728 - lw) / max(1, len(cols))
    body = theme.s["body"]
    label(s, theme, m, y + 48, lw, sd.get("corner", ""), color="inkSubtle", n=n)
    for i, c in enumerate(cols):
        x = m + lw + i * cw
        vline(s, x, y, y + 110, theme.col("line"))
        code_color = c.get("color", "accentText")
        if c.get("code"):
            text(s, theme, x + 32, y, cw - 48, 36, f"■ {c['code']}", theme.s["small"], color=code_color,
                 font=theme.mono, tracking=3, slide_no=n)
        text(s, theme, x + 32, y + 40, cw - 48, 60, c["name"], theme.s["h2"], color="ink", lh=1.1, slide_no=n)
    y += 110
    for r_i, r in enumerate(rows):
        hgt = max(est_height(cell, body, cw - 64, 1.45) for cell in r["cells"]) + 56
        hline(s, m, m + 1728, y, theme.col("lineStrong" if r_i == 0 else "line"))
        label(s, theme, m, y + 28, lw - 24, r["label"], n=n)
        for i, cell in enumerate(r["cells"]):
            x = m + lw + i * cw
            vline(s, x, y, y + hgt, theme.col("line"))
            text(s, theme, x + 32, y + 28, cw - 64, hgt - 40, cell, body, color="ink", lh=1.45, slide_no=n)
        y += hgt
    hline(s, m, m + 1728, y, theme.col("line"))
    check_bottom(theme, y, n, "matrix")
    return s


def L_table(prs, theme, meta, sd, n):
    """A real data table, drawn with hairlines. Keep to 8 rows or fewer."""
    s = new_slide(prs, theme)
    decor_grid(s, theme, n)
    chrome(s, theme, meta, page=n)
    m = theme.l["marginX"]
    y = header(s, theme, sd, n)
    cols = sd["columns"]
    widths = sd.get("widths") or [1 / len(cols)] * len(cols)
    tot = sum(widths)
    ws = [1728 * w / tot for w in widths]
    size = theme.s["small"]
    muted = set(sd.get("mutedColumns", []))
    rows = sd.get("rows", [])
    if len(rows) > 8:
        warn(n, f"table has {len(rows)} rows; split it or group rows (8 max)")
    hh = size * 1.4 + 24
    rect(s, m, y, 1728, hh, fill=theme.col("surfaceRaised"))
    x = m
    for i, c in enumerate(cols):
        text(s, theme, x + 20, y + 12, ws[i] - 40, size * 1.4, c, size, color="inkMuted", bold=True,
             **lab(theme), slide_no=n)
        x += ws[i]
    y += hh
    for r in rows:
        rh = max(est_height(cell, size, ws[i] - 40, 1.35) for i, cell in enumerate(r)) + 24
        rect(s, m, y, 1728, rh, fill=theme.col("surface"))
        hline(s, m, m + 1728, y, theme.col("line"))
        x = m
        for i, cell in enumerate(r):
            text(s, theme, x + 20, y + 12, ws[i] - 40, rh - 16, cell, size,
                 color="inkMuted" if i in muted else "ink", lh=1.35, slide_no=n)
            x += ws[i]
        y += rh
    hline(s, m, m + 1728, y, theme.col("line"))
    if sd.get("note"):
        text(s, theme, m, y + 28, 1728, 50, f"■ {sd['note']}", theme.s["body"], color="ink", lh=1.4, slide_no=n)
        y += 78
    check_bottom(theme, y, n, "table")
    return s


def L_tiers(prs, theme, meta, sd, n):
    """Ranked groups, each led by its count."""
    s = new_slide(prs, theme)
    decor_grid(s, theme, n)
    chrome(s, theme, meta, page=n)
    m = theme.l["marginX"]
    y = header(s, theme, sd, n)
    lw = 340
    for t_i, tier in enumerate(sd.get("tiers", [])):
        items = tier.get("items", [])
        per_row = 3 if len(items) > 1 else 1
        nrows = math.ceil(len(items) / per_row)
        ih = theme.s["body"] * 1.3 + theme.s["small"] * 1.3 + 4
        th = 22 + max(nrows * ih + (nrows - 1) * 14, theme.s["index"]) + 22
        hline(s, m, m + 1728, y, theme.col("lineStrong" if t_i == 0 else "line"))
        color = "accent" if t_i == 0 else ("ink" if t_i == 1 else "inkSubtle")
        cnt = tier.get("count", len(items))
        text(s, theme, m, y + 24, 120, theme.s["index"], str(cnt), theme.s["index"], color=color,
             font=theme.mono, lh=1.0, slide_no=n)
        text(s, theme, m + 130, y + 40, lw - 140, theme.s["small"] * 2.8, tier["label"], theme.s["small"],
             color="ink" if t_i < 2 else "inkMuted", bold=True, **lab(theme), lh=1.3, slide_no=n)
        cw = (1728 - lw) / per_row
        for i, it in enumerate(items):
            r, c = divmod(i, per_row)
            x = m + lw + c * cw
            yy = y + 22 + r * (ih + 14)
            text(s, theme, x, yy, cw - 24, theme.s["body"] * 1.3, it["name"], theme.s["body"], color="ink",
                 bold=True, slide_no=n)
            if it.get("desc"):
                text(s, theme, x, yy + theme.s["body"] * 1.3 + 4, cw - 24, theme.s["small"] * 1.3,
                     it["desc"], theme.s["small"], color="inkMuted", slide_no=n)
        y += th
    hline(s, m, m + 1728, y, theme.col("line"))
    check_bottom(theme, y, n, "tiers")
    return s


def L_layers(prs, theme, meta, sd, n):
    """Stacked architecture layers with one highlighted, plus a cross-cutting side note."""
    s = new_slide(prs, theme)
    decor_grid(s, theme, n)
    chrome(s, theme, meta, page=n)
    m = theme.l["marginX"]
    y = header(s, theme, sd, n)
    side = sd.get("side")
    lw = 1728 - (300 if side else 0)
    nw, namew = 96, 440
    dw = lw - nw - namew - 48
    start = y
    for i, ly in enumerate(sd.get("layers", [])):
        hi = ly.get("highlight")
        size = theme.s["body"]
        dh = est_height(ly.get("desc", ""), theme.s["small"], dw, 1.3)
        rh = max(dh, size * 1.3) + (48 if hi else 32)
        if hi:
            rect(s, m, y, lw, rh, fill=theme.col("surface"), line=theme.col("accent"), line_w=2)
        else:
            hline(s, m, m + lw, y, theme.col("lineStrong" if i == 0 else "line"))
        cy = y + rh / 2
        text(s, theme, m + 24, cy - 22, nw, 44, f"{i + 1:02d}", theme.s["h3"],
             color="accent" if hi else "inkSubtle", font=theme.mono, anchor="middle", check=False)
        text(s, theme, m + 24 + nw, cy - size * 0.7, namew, size * 1.4, ly["name"],
             size + (4 if hi else 0), color="ink", bold=True, anchor="middle", slide_no=n)
        text(s, theme, m + 24 + nw + namew, cy - dh / 2, dw, dh + 4, ly.get("desc", ""), theme.s["small"],
             color="ink" if hi else "inkMuted", lh=1.3, slide_no=n)
        y += rh
    hline(s, m, m + lw, y, theme.col("lineStrong"))
    if side:
        sx = m + lw + 40
        vline(s, sx, start, y, theme.col("inkSubtle"), 1, dash=True)
        items = side.get("items", [])
        bh = theme.s["small"] * 1.4 + 12 + len(items) * theme.s["body"] * 1.5
        sy = start + (y - start - bh) / 2
        label(s, theme, sx + 32, sy, 228, side.get("label", ""), n=n)
        sy += theme.s["small"] * 1.4 + 12
        for it in items:
            text(s, theme, sx + 32, sy, 228, theme.s["body"] * 1.5 * 2, it, theme.s["body"], color="ink",
                 lh=1.3, slide_no=n)
            sy += theme.s["body"] * 1.5 + (theme.s["body"] * 1.3 if len(it) > 16 else 0)
    check_bottom(theme, y, n, "layers")
    return s


def L_quadrant(prs, theme, meta, sd, n):
    """Four parallel ideas split by hairlines (no boxes)."""
    s = new_slide(prs, theme)
    decor_grid(s, theme, n)
    chrome(s, theme, meta, page=n)
    m = theme.l["marginX"]
    y = header(s, theme, sd, n)
    items = sd.get("items", [])
    if len(items) != 4:
        warn(n, "quadrant expects exactly 4 items; use matrix or tiers otherwise")
    hline(s, m, m + 1728, y, theme.col("lineStrong"))
    cw = 864
    cell_h = []
    for it in items:
        tw = cw - 48 - 104
        h = theme.s["h3"] * 1.25 + 12 + est_height(it.get("desc", ""), theme.s["body"], tw, 1.4) + 12 + \
            est_height(it.get("example", ""), theme.s["small"], tw, 1.4)
        cell_h.append(h + 72)
    rh = [max(cell_h[0:2] or [0]), max(cell_h[2:4] or [0])]
    vline(s, m + cw, y, y + sum(rh), theme.col("line"))
    hline(s, m, m + 1728, y + rh[0], theme.col("line"))
    for i, it in enumerate(items[:4]):
        r, c = divmod(i, 2)
        x = m + c * cw + (48 if c else 0)
        yy = y + (rh[0] if r else 0) + 36
        tw = cw - 48 - 104
        text(s, theme, x, yy, 80, theme.s["index"] * 0.9, str(i + 1), int(theme.s["index"] * 0.82),
             color="accent", font=theme.mono, lh=1.0, check=False)
        tx = x + 104
        text(s, theme, tx, yy, tw, theme.s["h3"] * 1.25, it["title"], theme.s["h3"], color="ink", bold=True,
             slide_no=n)
        yy += theme.s["h3"] * 1.25 + 12
        dh = est_height(it.get("desc", ""), theme.s["body"], tw, 1.4)
        text(s, theme, tx, yy, tw, dh, it.get("desc", ""), theme.s["body"], color="ink", lh=1.4, slide_no=n)
        yy += dh + 12
        if it.get("example"):
            eh = est_height(it["example"], theme.s["small"], tw, 1.4)
            text(s, theme, tx, yy, tw, eh, it["example"], theme.s["small"], color="inkMuted", italic=False,
                 lh=1.4, slide_no=n)
    check_bottom(theme, y + sum(rh), n, "quadrant")
    return s


def L_usecase(prs, theme, meta, sd, n):
    """Scenario: lead, triggers + stakeholders, numbered timeline, outcomes row."""
    s = new_slide(prs, theme)
    decor_grid(s, theme, n)
    chrome(s, theme, meta, page=n)
    m = theme.l["marginX"]
    y = header(s, theme, sd, n) - 8
    small, body = theme.s["small"], theme.s["body"]
    if sd.get("lead"):
        h = est_height(sd["lead"], theme.s["lead"], 1600, 1.45)
        text(s, theme, m, y, 1600, h, sd["lead"], theme.s["lead"], color="inkMuted", lh=1.45, slide_no=n)
        y += h + 28
    lw = 540
    ly = y
    label(s, theme, m, ly, lw, sd.get("triggersLabel", "Trigger"), n=n)
    ly += small * 1.4 + 12
    for t in sd.get("triggers", []):
        h = est_height(t, small, lw - 40, 1.35)
        text(s, theme, m, ly, lw - 40, h, t, small, color="ink", lh=1.35, slide_no=n)
        ly += h + 12
    if sd.get("stakeholders"):
        ly += 16
        label(s, theme, m, ly, lw, sd.get("stakeholdersLabel", "Stakeholder"), n=n)
        ly += small * 1.4 + 12
        h = est_height(sd["stakeholders"], small, lw - 40, 1.4)
        text(s, theme, m, ly, lw - 40, h, sd["stakeholders"], small, color="ink", lh=1.4, slide_no=n)
        ly += h
    rx = m + lw + 48
    rw = 1728 - lw - 48
    ry = y
    label(s, theme, rx, ry, rw, sd.get("stepsLabel", "Alur utama"), n=n)
    ry += small * 1.4 + 16
    steps = sd.get("steps", [])
    if len(steps) > 6:
        warn(n, f"{len(steps)} steps; condense to 5 or 6")
    for i, st in enumerate(steps):
        h = est_height(st, small, rw - 40, 1.35)
        style = "filled" if i in (0, len(steps) - 1) else "outline"
        square(s, rx, ry + 9, 14, theme, style)
        if i < len(steps) - 1:
            vline(s, rx + 7, ry + 30, ry + h + 22 + 9, theme.col("lineStrong"), 2)
        text(s, theme, rx + 38, ry, rw - 40, h, st, small, color="ink", lh=1.35, slide_no=n)
        ry += h + 22
    vline(s, m + lw, y, max(ly, y + 40), theme.col("line"))
    y = max(ly, ry) + 20
    outs = sd.get("outcomes", [])
    if outs:
        hline(s, m, m + 1728, y, theme.col("lineStrong"))
        ow = (1728 - 180) / len(outs)
        text(s, theme, m, y + 24, 170, small * 1.4, sd.get("outcomesLabel", "Outcome"), small, color="accentText",
             bold=True, **lab(theme), slide_no=n)
        oh = 0
        for i, o in enumerate(outs):
            x = m + 180 + i * ow
            if i:
                vline(s, x, y + 24, y + 24 + 70, theme.col("line"))
            h = est_height(o, body, ow - 64, 1.35)
            oh = max(oh, h)
            text(s, theme, x + 32, y + 24, ow - 64, h, o, body, color="ink", lh=1.35, slide_no=n)
        y += 24 + oh
    check_bottom(theme, y, n, "use case")
    return s


def L_phases(prs, theme, meta, sd, n):
    """Roadmap rail: phases with relative widths; the active phase is filled."""
    s = new_slide(prs, theme)
    decor_grid(s, theme, n)
    chrome(s, theme, meta, page=n)
    m = theme.l["marginX"]
    y = header(s, theme, sd, n) + 20
    phases = sd.get("phases", [])
    weights = [p.get("weight", 1) for p in phases]
    tot = sum(weights) or 1
    x = m
    bottom = y
    for i, ph in enumerate(phases):
        w = 1728 * weights[i] / tot
        active = ph.get("active", i == 0)
        if active:
            rect(s, x, y, 24, 24, fill=theme.col("accent"))
            rect(s, x + 24, y + 10, w - 24, 4, fill=theme.col("accent"))
        else:
            rect(s, x, y, 24, 24, fill=theme.col("ground"), line=theme.col("inkSubtle"), line_w=2)
            rect(s, x + 24, y + 11, w - 24, 2, fill=theme.col("lineStrong" if i == 1 else "line"))
        yy = y + 24 + 28
        pad = 64 if i < len(phases) - 1 else 0
        text(s, theme, x, yy, w - pad, theme.s["small"] * 1.4, ph.get("label", f"Phase {i + 1}"),
             theme.s["small"], color="accentText" if active else "inkMuted", font=theme.mono,
             **lab(theme), slide_no=n)
        yy += theme.s["small"] * 1.4 + 16
        th = est_height(ph["title"], theme.s["h2"], w - pad, 1.1)
        text(s, theme, x, yy, w - pad, th, ph["title"], theme.s["h2"], color="ink", lh=1.1, slide_no=n)
        yy += th + 16
        items = ph.get("items", [])
        two = len(items) > 4 and weights[i] >= 2
        colw = (w - pad) / (2 if two else 1)
        rows = math.ceil(len(items) / (2 if two else 1))
        for j, it in enumerate(items):
            r, c = (divmod(j, 2) if two else (j, 0))
            ih = est_height(it, theme.s["small"], colw - 24, 1.35) + 24
            iy = yy + r * (theme.s["small"] * 1.35 * 2 + 24 if two else ih)
            hline(s, x + c * colw, x + (c + 1) * colw - (24 if two and c == 0 else 0), iy, theme.col("line"))
            text(s, theme, x + c * colw, iy + 12, colw - 24, ih - 16, it, theme.s["small"], color="ink",
                 lh=1.35, slide_no=n)
        last = yy + (rows * (theme.s["small"] * 1.35 * 2 + 24) if two else
                     sum(est_height(it, theme.s["small"], colw - 24, 1.35) + 24 for it in items))
        bottom = max(bottom, last)
        x += w
    check_bottom(theme, bottom, n, "phases")
    return s


def L_coverage(prs, theme, meta, sd, n):
    """Progress across many items: legend, segmented bar, grouped columns."""
    s = new_slide(prs, theme)
    decor_grid(s, theme, n)
    chrome(s, theme, meta, page=n)
    m = theme.l["marginX"]
    y = header(s, theme, sd, n)
    statuses = sd.get("statuses", [
        {"key": "done", "label": "Done", "style": "filled"},
        {"key": "next", "label": "Next", "style": "outline"},
        {"key": "planned", "label": "Planned", "style": "muted"}])
    groups = sd.get("groups", [])
    counts = {st["key"]: 0 for st in statuses}
    for g in groups:
        for it in g["items"]:
            counts[it.get("status", statuses[-1]["key"])] = counts.get(it.get("status"), 0) + 1
    small = theme.s["small"]
    x = m
    for st in statuses:
        square(s, x, y + 9, 16, theme, st["style"])
        lbl = f"{st['label']}  ·  {counts.get(st['key'], 0)}"
        w = len(lbl) * small * 0.55 + 40
        text(s, theme, x + 28, y, w, small * 1.4, lbl, small, color="ink", slide_no=n)
        x += 28 + w + 24
    y += small * 1.4 + 16
    total = sum(counts.values()) or 1
    gap = 6
    cell = (1728 - gap * (total - 1)) / total
    i = 0
    for st in statuses:
        for _ in range(counts.get(st["key"], 0)):
            cx = m + i * (cell + gap)
            if st["style"] == "filled":
                rect(s, cx, y, cell, 28, fill=theme.col("accent"))
            elif st["style"] == "outline":
                rect(s, cx, y, cell, 28, line=theme.col("accentText"), line_w=2)
            else:
                rect(s, cx, y, cell, 28, fill=theme.col("line"))
            i += 1
    y += 28 + 40
    hline(s, m, m + 1728, y, theme.col("lineStrong"))
    gw = 1728 / max(1, len(groups))
    style_of = {st["key"]: st["style"] for st in statuses}
    deepest = y
    up = lab(theme)["upper"]
    head_lines = 1
    for g_i, g in enumerate(groups):
        inner = gw - (24 if g_i == 0 else 48)
        head_lines = max(head_lines, est_lines(g["name"], small, inner - 70, upper=up,
                                               tracking=lab(theme)["tracking"] * 0.5))
    head_h = head_lines * small * 1.25
    for g_i, g in enumerate(groups):
        gx = m + g_i * gw
        if g_i:
            vline(s, gx, y, y + 40 + len(max(groups, key=lambda q: len(q["items"]))["items"]) * (small * 1.3 + 14) + 50,
                  theme.col("line"))
        px0 = gx + (0 if g_i == 0 else 24)
        inner = gw - (24 if g_i == 0 else 48)
        done = sum(1 for it in g["items"] if style_of.get(it.get("status")) == "filled")
        text(s, theme, px0, y + 24, inner - 70, head_h + 4, g["name"], small, color="inkMuted", bold=True,
             **lab(theme), lh=1.2, slide_no=n)
        text(s, theme, px0 + inner - 60, y + 24, 60, small * 1.4, f"{done}/{len(g['items'])}", small,
             color="ink", font=theme.mono, align="right", check=False)
        yy = y + 24 + head_h + 20
        for it in g["items"]:
            stl = style_of.get(it.get("status"), "muted")
            square(s, px0, yy + 8, 16, theme, stl)
            h = est_height(it["name"], small, inner - 30, 1.3)
            text(s, theme, px0 + 30, yy, inner - 30, h, it["name"], small,
                 color="ink" if stl != "muted" else "inkSubtle", lh=1.3, slide_no=n)
            yy += h + 14
        deepest = max(deepest, yy)
    check_bottom(theme, deepest, n, "coverage groups")
    return s


def L_evidence(prs, theme, meta, sd, n):
    """1-3 photos with captions, numbered findings on the right."""
    s = new_slide(prs, theme)
    decor_grid(s, theme, n)
    chrome(s, theme, meta, page=n)
    m = theme.l["marginX"]
    y = header(s, theme, sd, n)
    photos = sd.get("photos", [])[:3]
    pw = 540 if len(photos) <= 2 else 360
    ph = pw * 0.75
    x = m
    for p in photos:
        if not image(s, p.get("path"), x, y, pw, ph):
            rect(s, x, y, pw, ph, fill=theme.col("surface"), line=theme.col("line"))
            text(s, theme, x, y, pw, ph, "[photo]", theme.s["small"], color="inkSubtle", align="center",
                 anchor="middle", check=False)
        text(s, theme, x, y + ph + 14, pw, theme.s["small"] * 1.35, p.get("caption", ""), theme.s["small"],
             color="ink", lh=1.35, slide_no=n)
        if p.get("sub"):
            text(s, theme, x, y + ph + 14 + theme.s["small"] * 1.35, pw, theme.s["small"] * 1.35, p["sub"],
                 theme.s["small"], color="inkMuted", lh=1.35, slide_no=n)
        x += pw + 24
    fx = x + 32
    fw = m + 1728 - fx
    fy = y
    label(s, theme, fx, fy, fw, sd.get("findingsLabel", "Findings"), n=n)
    fy += theme.s["small"] * 1.4 + 14
    for i, f in enumerate(sd.get("findings", [])):
        h = est_height(f, theme.s["small"], fw - 44, 1.4) + 32
        hline(s, fx, fx + fw, fy, theme.col("lineStrong" if i == 0 else "line"))
        text(s, theme, fx, fy + 16, 30, 34, str(i + 1), theme.s["small"], color="accentText", font=theme.mono,
             check=False)
        text(s, theme, fx + 44, fy + 16, fw - 44, h - 24, f, theme.s["small"], color="ink", lh=1.4, slide_no=n)
        fy += h
    hline(s, fx, fx + fw, fy, theme.col("line"))
    check_bottom(theme, max(fy, y + ph + 80), n, "evidence")
    return s


def L_decisions(prs, theme, meta, sd, n):
    """Asks for the audience: headline + framing on the left, numbered asks right."""
    s = new_slide(prs, theme)
    decor_grid(s, theme, n)
    chrome(s, theme, meta, page=n)
    m, top = theme.l["marginX"], theme.l["marginTop"]
    lw = 620
    y = top
    if sd.get("eyebrow"):
        text(s, theme, m, y, lw, theme.s["small"] * 1.4, sd["eyebrow"], theme.s["small"], color="accentText",
             bold=True, **eyebrow_style(theme), slide_no=n)
        y += theme.s["small"] * 1.4 + 12
    th = est_height(sd.get("title", ""), theme.s["h1"], lw, 1.15)
    text(s, theme, m, y, lw, th, sd.get("title", ""), theme.s["h1"], color="ink", lh=1.15, slide_no=n)
    if sd.get("lead"):
        lh = est_height(sd["lead"], theme.s["lead"], lw, 1.5)
        text(s, theme, m, bottom_limit(theme) - lh, lw, lh, sd["lead"], theme.s["lead"], color="inkMuted",
             lh=1.5, slide_no=n)
    rx = m + lw + 96
    rw = 1728 - lw - 96
    items = sd.get("items", [])
    hs = []
    for it in items:
        hs.append(40 + theme.s["h2"] * 1.2 + 10 + est_height(it.get("desc", ""), theme.s["body"], rw - 136, 1.45) + 40)
    ry = top + (bottom_limit(theme) - top - sum(hs)) / 2
    for i, it in enumerate(items):
        hline(s, rx, rx + rw, ry, theme.col("lineStrong" if i == 0 else "line"))
        text(s, theme, rx, ry + 40, 96, theme.s["metric"] * 0.6, str(i + 1), int(theme.s["metric"] * 0.6),
             color="accent", font=theme.mono, lh=1.0, check=False)
        text(s, theme, rx + 136, ry + 40, rw - 136, theme.s["h2"] * 1.2, it["title"], theme.s["h2"], color="ink",
             slide_no=n)
        dh = est_height(it.get("desc", ""), theme.s["body"], rw - 136, 1.45)
        text(s, theme, rx + 136, ry + 40 + theme.s["h2"] * 1.2 + 10, rw - 136, dh, it.get("desc", ""),
             theme.s["body"], color="inkMuted", lh=1.45, slide_no=n)
        ry += hs[i]
    hline(s, rx, rx + rw, ry, theme.col("line"))
    check_bottom(theme, ry, n, "decisions")
    return s


LAYOUTS = {
    "cover": L_cover, "closing": L_closing, "divider": L_divider, "summary": L_summary,
    "statement": L_statement, "convergence": L_convergence, "matrix": L_matrix, "table": L_table,
    "tiers": L_tiers, "layers": L_layers, "quadrant": L_quadrant, "usecase": L_usecase,
    "phases": L_phases, "coverage": L_coverage, "evidence": L_evidence, "decisions": L_decisions,
}


# --------------------------------------------------------------------------- notes + copy lint

BANNED = ["unlock", "seamless", "elevate", "empower", "game-changer", "game changer", "revolutionary",
          "cutting-edge", "cutting edge", "next-level", "world-class", "industry-leading", "synergy",
          "leverage", "robust", "delve", "testament", "journey", "landscape", "paradigm",
          "revolusioner", "terdepan", "kelas dunia", "mutakhir"]


def notes_text(notes, lang):
    if not notes:
        return ""
    if isinstance(notes, str):
        return notes
    key_lbl, tr_lbl = ("PESAN UTAMA", "TRANSISI") if lang == "id" else ("KEY MESSAGE", "TRANSITION")
    out = []
    if notes.get("key"):
        out.append(f"{key_lbl}: {notes['key']}")
        out.append("")
    for p in notes.get("points", []):
        out.append(f"• {p}")
    if notes.get("transition"):
        out.append("")
        out.append(f"{tr_lbl}: {notes['transition']}")
    return "\n".join(out)


def walk_strings(obj):
    if isinstance(obj, str):
        yield obj
    elif isinstance(obj, dict):
        for k, v in obj.items():
            if k in ("image", "path", "logoLeft", "logoRight", "logoCover", "layout", "color", "status",
                     "style", "key", "preset"):
                continue
            yield from walk_strings(v)
    elif isinstance(obj, list):
        for v in obj:
            yield from walk_strings(v)


def lint_copy(deck):
    for i, sd in enumerate(deck["slides"], 1):
        for sval in walk_strings(sd):
            if "\u2014" in sval or "\u2013" in sval or " -- " in sval:
                warn(i, f"em/en dash found, rewrite with a period, comma, colon or parentheses: \"{sval[:50]}\"")
            low = sval.lower()
            for w in BANNED:
                if re.search(r"\b" + re.escape(w) + r"\b", low):
                    warn(i, f"empty buzzword \"{w}\": say the concrete thing instead")
            if re.search(r"\b(it's not|bukan hanya|bukan sekadar)\b.*\b(it's|tetapi|tapi|melainkan)\b", low):
                warn(i, "\"not X, but Y\" construction; state the point directly")
        if sd.get("layout") not in ("cover", "closing", "divider") and not sd.get("notes"):
            warn(i, "no speaker notes / talking points")
    for sval in walk_strings(deck.get("meta", {})):
        if "\u2014" in sval or "\u2013" in sval:
            WARNINGS.append(f"meta: em/en dash in \"{sval[:50]}\"")


# --------------------------------------------------------------------------- main

def build(deck, theme, out_path, deck_dir):
    def resolve(pth):
        if pth and not os.path.isabs(pth):
            return os.path.normpath(os.path.join(deck_dir, pth))
        return pth

    meta = dict(deck.get("meta", {}))
    for k in ("logoLeft", "logoRight", "logoCover"):
        if meta.get(k):
            meta[k] = resolve(meta[k])
    prs = Presentation()
    prs.slide_width, prs.slide_height = px(CANVAS_W), px(CANVAS_H)
    lang = meta.get("language", "en")
    for n, sd in enumerate(deck["slides"], 1):
        sd = json.loads(json.dumps(sd))
        if sd.get("image"):
            sd["image"] = resolve(sd["image"])
        for p in sd.get("photos", []):
            p["path"] = resolve(p.get("path"))
        fn = LAYOUTS.get(sd.get("layout"))
        if not fn:
            warn(n, f"unknown layout \"{sd.get('layout')}\"; known: {', '.join(LAYOUTS)}")
            continue
        slide = fn(prs, theme, meta, sd, n)
        nt = notes_text(sd.get("notes"), lang)
        if nt:
            slide.notes_slide.notes_text_frame.text = nt
    if out_path:
        prs.save(out_path)
    return prs


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("deck", help="deck spec JSON")
    ap.add_argument("-t", "--tokens", help="token preset file or name (default: deck.meta.preset)")
    ap.add_argument("-o", "--out", help="output .pptx path")
    ap.add_argument("--check", action="store_true", help="lint only; do not write a file")
    a = ap.parse_args()
    with open(a.deck, encoding="utf-8") as fh:
        deck = json.load(fh)
    theme = load_tokens(a.tokens or deck.get("meta", {}).get("preset", "neutral-light"))
    lint_copy(deck)
    build(deck, theme, None if a.check else (a.out or os.path.splitext(a.deck)[0] + ".pptx"),
          os.path.dirname(os.path.abspath(a.deck)))
    if WARNINGS:
        print(f"{len(WARNINGS)} warning(s):")
        for w in WARNINGS:
            print("  - " + w)
    else:
        print("no warnings")
    if not a.check:
        print("wrote", a.out or os.path.splitext(a.deck)[0] + ".pptx")


if __name__ == "__main__":
    main()
